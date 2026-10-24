"""Render verified source references in editable exports and offline appendices."""

from __future__ import annotations

import html
import json
import re

MARKER = re.compile(r"\[\^([a-zA-Z0-9_-]+)\]")


def detail(row):
    label, excerpt = row["tool"], row["excerpt"]
    try:
        data = json.loads(excerpt)
        record = data.get("record", data) if isinstance(data, dict) else {}
        locator = str(record.get("locator", data.get("locator", "")))
        filename = str(record.get("filename", record.get("name", "")))
        if filename or locator:
            label = " · ".join(value for value in (filename, locator) if value)
        if "value" in record:
            excerpt = str(record["value"])
            if record.get("formula"):
                excerpt += (
                    " (formula; saved cached value: "
                    + str(record.get("cached_value"))
                    + "; not evaluated by OpenDuet)"
                )
        else:
            excerpt = json.dumps(record, ensure_ascii=False)
    except (ValueError, AttributeError, TypeError):
        pass
    return {
        "id": row["id"],
        "tool": row["tool"],
        "label": label[:500],
        "url": row["url"],
        "excerpt": excerpt[:1200],
    }


def value_text(value):
    if isinstance(value, dict) and "value" in value:
        return with_ids(str(value["value"]), value)
    return str(value)


def with_ids(text, item):
    return str(text) + "".join(" [^" + ident + "]" for ident in item.get("source_ids", []))


def parts(text):
    cursor = 0
    for match in MARKER.finditer(str(text)):
        yield str(text)[cursor : match.start()], None
        yield match[0], match[1]
        cursor = match.end()
    yield str(text)[cursor:], None


def doc_text(paragraph, text, sources):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for value, ident in parts(text):
        if ident is None:
            paragraph.add_run(value)
            continue
        source = sources[ident]
        link = OxmlElement("w:hyperlink")
        if source["url"]:
            link.set(qn("r:id"), paragraph.part.relate_to(source["url"], RT.HYPERLINK, is_external=True))
        else:
            link.set(qn("w:anchor"), "src_" + ident)
        run, props, style, node = (OxmlElement(tag) for tag in ("w:r", "w:rPr", "w:rStyle", "w:t"))
        style.set(qn("w:val"), "Hyperlink")
        props.append(style)
        node.text = value
        run.extend([props, node])
        link.append(run)
        paragraph._p.append(link)


def doc_appendix(document, sources):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    if sources:
        document.add_heading("Sources", 1)
    for index, source in enumerate(sources.values(), 1):
        paragraph = document.add_paragraph()
        start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        start.set(qn("w:id"), str(index))
        start.set(qn("w:name"), "src_" + source["id"])
        end.set(qn("w:id"), str(index))
        paragraph._p.append(start)
        paragraph.add_run(source["label"] + " " + source["url"] + " ")
        doc_text(paragraph, "[^" + source["id"] + "]", sources)
        paragraph._p.append(end)
        document.add_paragraph(source["excerpt"])


def ppt_text(frame, text, sources, links):
    frame.clear()
    for index, line in enumerate(str(text).splitlines() or [""]):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        for value, ident in parts(line):
            run = paragraph.add_run()
            run.text = value
            if ident is not None:
                if sources[ident]["url"]:
                    run.hyperlink.address = sources[ident]["url"]
                else:
                    links.append((run, ident))


def ppt_appendix(deck, sources, links):
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt

    rows, targets, previews = list(sources.values()), {}, []
    for index in range(0, len(rows), 4):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(12), Inches(0.8))
        title.text_frame.text = "Sources"
        title.text_frame.paragraphs[0].font.size = Pt(30)
        body = []
        for offset, source in enumerate(rows[index : index + 4]):
            targets[source["id"]] = slide
            box = slide.shapes.add_textbox(Inches(0.7), Inches(1.5 + offset * 1.3), Inches(12), Inches(1.2))
            box.text_frame.word_wrap = True
            text = source["label"] + " [^" + source["id"] + "]\n" + source["excerpt"][:300]
            ppt_text(box.text_frame, text, sources, links)
            for paragraph in box.text_frame.paragraphs:
                paragraph.font.size = Pt(13)
            body.append(text)
        previews.append(
            {
                "slide": len(deck.slides),
                "title": "Sources",
                "layout": "bullets",
                "body": "\n".join(body),
                "notes": "",
                "needs_review": False,
            }
        )
    for run, ident in links:
        target = targets[ident]
        link = OxmlElement("a:hlinkClick")
        link.set(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id",
            run.part.relate_to(target.part, RT.SLIDE),
        )
        link.set("action", "ppaction://hlinksldjump")
        run._r.get_or_add_rPr().append(link)
    return previews


def pdf_text(text, sources):
    result = []
    for value, ident in parts(text):
        if ident is None:
            result.append(html.escape(value))
        else:
            href = sources[ident]["url"] or "#src_" + ident
            result.append(
                '<link href="'
                + html.escape(href, quote=True)
                + '" color="#31522b">'
                + html.escape(value)
                + "</link>"
            )
    return "".join(result)
