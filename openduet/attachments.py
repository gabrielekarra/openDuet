"""Bounded, inert task attachments. Never execute or serve uploaded HTML."""

from __future__ import annotations

import asyncio
import io
import zipfile
from pathlib import Path

from .store import uid

TYPES = {
    ".txt": "text/plain",
    ".md": "text/plain",
    ".csv": "text/csv",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".ogg": "audio/ogg",
    ".mp3": "audio/mpeg",
    ".wav": "audio/wav",
    ".m4a": "audio/mp4",
}


async def ingest(runtime, data, name, task=None, *, inbound=True):
    suffix = Path(name).suffix.lower()
    if suffix not in TYPES or len(data) > 20_000_000:
        raise ValueError("Unsupported attachment type or size (maximum 20 MB)")
    directory = runtime.config.data_dir / "attachments"
    directory.mkdir(exist_ok=True, mode=0o700)
    ident = uid()
    path = directory / (ident + suffix)
    text, transcript = "", ""
    if suffix in {".txt", ".md", ".csv"}:
        text = data[:100000].decode("utf-8", errors="replace")
    elif suffix in {".docx", ".xlsx", ".pptx"}:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 40_000_000:
                raise ValueError("Expanded attachment too large")
        if suffix == ".docx":
            from docx import Document

            text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
        elif suffix == ".pptx":
            from pptx import Presentation

            deck = Presentation(io.BytesIO(data))
            text = "\n".join(
                shape.text
                for slide in list(deck.slides)[:50]
                for shape in slide.shapes
                if shape.has_text_frame
            )[:32000]
        else:
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            rows = []
            for sheet in workbook:
                for index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if index >= 1000:
                        break
                    rows.append(sheet.title + ": " + str(row[:30]))
            text = "\n".join(rows)
            workbook.close()
    elif suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text()[:20000] for page in list(reader.pages)[:40])
    elif suffix in {".png", ".jpg", ".jpeg"}:
        from PIL import Image

        with Image.open(io.BytesIO(data)) as picture:
            if picture.width * picture.height > 20_000_000:
                raise ValueError("Image exceeds the 20 megapixel decode limit")
            picture.verify()
    path.write_bytes(data)
    path.chmod(0o600)
    if TYPES[suffix].startswith("audio/"):
        if not runtime.config.voice_model:
            text = "Voice input needs a local Whisper model. Configure voice_model in Settings."
        else:

            def transcribe():
                from faster_whisper import WhisperModel

                model = WhisperModel(
                    runtime.config.voice_model, device="cpu", compute_type="int8", local_files_only=True
                )
                segments, info = model.transcribe(str(path), vad_filter=True)
                if info.duration > 180:
                    raise ValueError("Voice input is limited to three minutes")
                return " ".join(segment.text for segment in segments)[:32000]

            transcript = await asyncio.to_thread(transcribe)
            text = transcript
    runtime.store.execute(
        "INSERT INTO attachments VALUES(?,?,?,?,?,?)",
        (ident, task, Path(name).name[:200], TYPES[suffix], path.name, len(data)),
    )
    if inbound:
        runtime.store.execute("INSERT INTO attachment_inputs VALUES(?)", (ident,))
    return {
        "attachment": ident,
        "name": Path(name).name[:200],
        "text": text[:32000],
        "transcript": transcript,
    }
