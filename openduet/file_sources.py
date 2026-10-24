"""Inert, bounded file snapshots with precise locators and task-scoped reads."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import zipfile
from pathlib import Path

MAX_RECORDS, MAX_TEXT = 2000, 200000


def ensure(store):
    store.execute(
        "CREATE TABLE IF NOT EXISTS attachment_extracts(id TEXT PRIMARY KEY,sha256 TEXT,records TEXT,truncated INTEGER)"
    )
    store.execute("CREATE TABLE IF NOT EXISTS task_attachment_access(task TEXT,id TEXT,PRIMARY KEY(task,id))")


def extract(data, name):
    suffix = Path(name).suffix.lower()
    records, size, clipped = [], 0, False

    def add(locator, value, **metadata):
        nonlocal size, clipped
        if value is None or value == "":
            return True
        if len(records) >= MAX_RECORDS or size >= MAX_TEXT:
            clipped = True
            return False
        text = str(value)
        record = {
            "locator": locator,
            "value": value if isinstance(value, (int, float, bool)) else text[:4000],
            **metadata,
        }
        if len(text) > 4000:
            record["truncated"] = clipped = True
        if isinstance(record.get("cached_value"), str) and len(record["cached_value"]) > 4000:
            record["cached_value"] = record["cached_value"][:4000]
            record["truncated"] = clipped = True
        encoded_size = len(json.dumps(record, default=str))
        if size + encoded_size > MAX_TEXT:
            clipped = True
            return False
        records.append(record)
        size += encoded_size
        return True

    if suffix in {".txt", ".md"}:
        for index, line in enumerate(data.decode("utf-8-sig", errors="replace").splitlines(), 1):
            if not add(f"line {index}", line):
                break
    elif suffix == ".csv":
        from openpyxl.utils import get_column_letter

        for index, row in enumerate(csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace"))), 1):
            if index > 1000:
                clipped = True
                break
            for column, value in enumerate(row[:30], 1):
                if not add(f"CSV!{get_column_letter(column)}{index}", value):
                    break
            if len(row) > 30:
                clipped = True
            if len(records) >= MAX_RECORDS or size >= MAX_TEXT:
                clipped = True
                break
    elif suffix in {".docx", ".xlsx", ".pptx"}:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 40_000_000:
                raise ValueError("Expanded attachment too large")
        if suffix == ".docx":
            from docx import Document

            document = Document(io.BytesIO(data))
            for index, paragraph in enumerate(document.paragraphs, 1):
                if not add(f"paragraph {index}", paragraph.text):
                    break
            for index, table in enumerate(document.tables[:100], 1):
                for row_index, row in enumerate(table.rows[:100], 1):
                    for column, cell in enumerate(row.cells[:30], 1):
                        if not add(f"table {index}!R{row_index}C{column}", cell.text):
                            break
                    if len(records) >= MAX_RECORDS or size >= MAX_TEXT:
                        break
                if len(records) >= MAX_RECORDS or size >= MAX_TEXT:
                    break
        elif suffix == ".pptx":
            from pptx import Presentation

            deck = Presentation(io.BytesIO(data))
            clipped = len(deck.slides) > 50
            for index, slide in enumerate(list(deck.slides)[:50], 1):
                for shape_index, shape in enumerate(slide.shapes, 1):
                    if shape.has_text_frame and not add(f"slide {index}, shape {shape_index}", shape.text):
                        break
                    if shape.has_table:
                        for r, row in enumerate(list(shape.table.rows)[:100], 1):
                            for c, cell in enumerate(row.cells[:30], 1):
                                if not add(f"slide {index}, table {shape_index}!R{r}C{c}", cell.text):
                                    break
                if len(records) >= MAX_RECORDS or size >= MAX_TEXT:
                    break
        else:
            from openpyxl import load_workbook

            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in book:
                    max_row, max_column = min(sheet.max_row or 1000, 1000), min(sheet.max_column or 30, 30)
                    clipped |= (sheet.max_row or 0) > 1000 or (sheet.max_column or 0) > 30
                    values = cached[sheet.title].iter_rows(max_row=max_row, max_col=max_column)
                    for row, saved in zip(
                        sheet.iter_rows(max_row=max_row, max_col=max_column), values, strict=True
                    ):
                        for cell, cache in zip(row, saved, strict=True):
                            if cell.value is None:
                                continue
                            metadata = (
                                {"formula": True, "cached_value": cache.value, "evaluated_by_host": False}
                                if cell.data_type == "f"
                                else {}
                            )
                            if not add(
                                f"'{sheet.title.replace(chr(39), chr(39) * 2)}'!{cell.coordinate}",
                                cell.value,
                                **metadata,
                            ):
                                break
                        if len(records) >= MAX_RECORDS or size >= MAX_TEXT:
                            break
                    if len(records) >= MAX_RECORDS or size >= MAX_TEXT:
                        break
            finally:
                book.close()
                cached.close()
    elif suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        clipped = len(reader.pages) > 40
        for index, page in enumerate(list(reader.pages)[:40], 1):
            if not add(f"page {index}", page.extract_text() or ""):
                break
    return records, clipped or len(records) >= MAX_RECORDS or size >= MAX_TEXT


def save(store, ident, data, records, truncated):
    ensure(store)
    store.execute(
        "INSERT OR REPLACE INTO attachment_extracts VALUES(?,?,?,?)",
        (
            ident,
            hashlib.sha256(data).hexdigest(),
            json.dumps(records, ensure_ascii=False, default=str),
            int(truncated),
        ),
    )


def attach(store, task, identifiers, db):
    for ident in dict.fromkeys(identifiers):
        row = db.execute(
            "SELECT a.id,a.task FROM attachments a JOIN attachment_inputs i ON i.id=a.id WHERE a.id=?",
            (ident,),
        ).fetchone()
        if not row:
            raise ValueError("Input attachment does not exist")
        if row["task"] is None:
            db.execute("UPDATE attachments SET task=? WHERE id=?", (task, ident))
        else:
            db.execute("INSERT OR IGNORE INTO task_attachment_access VALUES(?,?)", (task, ident))


def snapshot(runtime, row):
    rows = runtime.store.rows("SELECT * FROM attachment_extracts WHERE id=?", (row["id"],))
    data = (runtime.config.data_dir / "attachments" / row["filename"]).read_bytes()
    if len(data) > 20_000_000:
        raise ValueError("Attachment exceeds the 20 MB limit")
    if not rows:
        records, truncated = extract(data, row["name"])
        save(runtime.store, row["id"], data, records, truncated)
        rows = runtime.store.rows("SELECT * FROM attachment_extracts WHERE id=?", (row["id"],))
    result = rows[0]
    if hashlib.sha256(data).hexdigest() != result["sha256"]:
        raise ValueError("Attachment changed since import; upload a new copy")
    return result


def listing(runtime, task):
    result = []
    for row in runtime.store.rows(
        "SELECT a.* FROM attachments a JOIN attachment_inputs i ON i.id=a.id WHERE a.task=? OR EXISTS(SELECT 1 FROM task_attachment_access p WHERE p.task=? AND p.id=a.id) LIMIT 20",
        (task, task),
    ):
        cached = snapshot(runtime, row)
        result.append(
            {
                "attachment": row["id"],
                "name": row["name"],
                "records": len(json.loads(cached["records"])),
                "truncated": bool(cached["truncated"]),
                "sha256": cached["sha256"],
            }
        )
    return result


def read(runtime, task, attachment, locator="", offset=0, limit=20):
    rows = runtime.store.rows(
        "SELECT a.* FROM attachments a JOIN attachment_inputs i ON i.id=a.id WHERE a.id=? AND (a.task=? OR EXISTS(SELECT 1 FROM task_attachment_access p WHERE p.task=? AND p.id=a.id))",
        (attachment, task, task),
    )
    if not rows:
        raise ValueError("Only input files attached to this task can be read")
    row = rows[0]
    cached = snapshot(runtime, row)
    records = json.loads(cached["records"])
    if locator:
        records = [record for record in records if record["locator"] == locator]
        if not records:
            raise ValueError("Locator not present in this file's extracted snapshot")
    if not isinstance(offset, int) or not isinstance(limit, int) or offset < 0 or not 1 <= limit <= 20:
        raise ValueError("Use a nonnegative offset and a limit between 1 and 20")
    selected, length = [], 0
    for record in records[offset : offset + limit]:
        length += len(json.dumps(record))
        if selected and length > 24000:
            break
        selected.append(record)
    page = [
        {
            "id": attachment + ":" + record["locator"],
            "attachment": attachment,
            "filename": row["name"],
            **record,
        }
        for record in selected
    ]
    return {
        "attachment": attachment,
        "filename": row["name"],
        "sha256": cached["sha256"],
        "items": page,
        "next_offset": offset + len(page) if offset + len(page) < len(records) else None,
        "truncated": bool(cached["truncated"]),
    }
