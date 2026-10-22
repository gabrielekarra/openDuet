from __future__ import annotations

import csv
import io
import time
from pathlib import Path

from .store import uid

FORMATS = {"md", "txt", "html", "csv", "xlsx", "docx", "pptx", "dashboard", "webapp", "pdf"}


def _plain_create(store, directory: Path, task, title, format, content):
    if format not in FORMATS:
        raise ValueError("Unsupported artifact format")
    directory.mkdir(mode=0o700, parents=True, exist_ok=True)
    ident = uid()
    filename = ident + "." + format
    path = directory / filename
    if format in {"md", "txt", "html", "csv"}:
        path.write_text(content, encoding="utf-8")
    elif format == "xlsx":
        from openpyxl import Workbook

        book = Workbook()
        for row in csv.reader(io.StringIO(content)):
            # Spreadsheet formulas can trigger external calls; generated cells are literal data.
            book.active.append(["'" + c if c.startswith(("=", "+", "-", "@")) else c for c in row])
        book.save(path)
    elif format == "docx":
        from docx import Document

        doc = Document()
        doc.add_heading(title, 0)
        for line in content.splitlines():
            doc.add_paragraph(line)
        doc.save(path)
    elif format == "pptx":
        from pptx import Presentation

        deck = Presentation()
        for block in content.split("\n---\n"):
            lines = block.strip().splitlines() or [title]
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = lines[0].lstrip("# ")
            slide.placeholders[1].text = "\n".join(lines[1:])
        deck.save(path)
    path.chmod(0o600)
    store.execute(
        "INSERT INTO artifacts VALUES(?,?,?,?,?,?)", (ident, task, title, format, filename, time.time())
    )
    return {"id": ident, "title": title, "format": format, "download": f"/api/artifacts/{ident}"}


def create(store, directory: Path, task, title, format, content, parent=None, refresh=None):
    from .make import publish

    return publish(store, directory, task, title, format, content, parent, refresh, _plain_create)
