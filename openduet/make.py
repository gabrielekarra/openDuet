"""Versioned deliverables, typed Office files and declarative dashboards."""

from __future__ import annotations

import csv
import difflib
import html
import io
import json
import re
import shutil
import subprocess
import tempfile
import time
import zipfile
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from . import provenance
from .store import uid


def ensure(store):
    with store.db() as db:
        db.executescript("""
            CREATE TABLE IF NOT EXISTS artifact_versions(
                id TEXT PRIMARY KEY,root TEXT,version INTEGER,parent TEXT,spec TEXT,metadata TEXT);
            CREATE TABLE IF NOT EXISTS dashboard_refresh(
                root TEXT PRIMARY KEY,local_time TEXT,timezone TEXT,prompt TEXT,provider TEXT);
            CREATE TABLE IF NOT EXISTS dashboard_runs(root TEXT,local_day TEXT,task TEXT,PRIMARY KEY(root,local_day));
        """)


def validate_formula(value):
    if not isinstance(value, str) or not value.startswith("=") or len(value) > 2000:
        raise ValueError("A formula must start with = and contain at most 2000 characters")
    if re.search(
        r"\[|\]|https?://|file:|\||(?:WEBSERVICE|HYPERLINK|IMPORTXML|IMPORTDATA|DDE|RTD|CALL|EXEC|REGISTER|INDIRECT)\s*\(",
        value,
        re.I,
    ):
        raise ValueError("External, executable and indirect spreadsheet formulas are not supported")
    return value


def workbook(path, spec, sources=None, source_sheet="Sources"):
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.comments import Comment
    from openpyxl.workbook.properties import CalcProperties

    book = Workbook()
    book.remove(book.active)
    sheets = spec.get("sheets", [])
    if not 1 <= len(sheets) <= 30:
        raise ValueError("Workbook requires 1–30 sheets")
    for sheet in sheets:
        page = book.create_sheet(str(sheet["name"])[:31])
        cells = sheet.get("cells", {})
        if len(cells) > 10000:
            raise ValueError("Too many workbook cells")
        for address, value in cells.items():
            if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,5}", address):
                raise ValueError("Invalid cell reference")
            target = page[address]
            if isinstance(value, dict):
                if "formula" in value:
                    target.value = validate_formula(value["formula"])
                else:
                    target.value = value.get("value")
                    if isinstance(target.value, str) and target.value.startswith(("=", "+", "@")):
                        target.value = "'" + target.value
                target.number_format = str(value.get("number_format", "General"))[:100]
                if ids := value.get("source_ids"):
                    rows = [(sources or {})[ident] for ident in ids]
                    target.comment = Comment(
                        "\n".join(
                            f"[^{row['id']}] {provenance.detail(row)['label']} {row['url']}\n{provenance.detail(row)['excerpt']}"
                            for row in rows
                        ),
                        "OpenDuet sources",
                    )
                    if rows[0]["url"]:
                        target.hyperlink = rows[0]["url"]
                    else:
                        from openpyxl.worksheet.hyperlink import Hyperlink

                        target.hyperlink = Hyperlink(
                            ref=address,
                            location="'" + source_sheet + "'!A" + str(list(sources).index(rows[0]["id"]) + 2),
                        )
            elif isinstance(value, str) and value.startswith(("=", "+", "@")):
                target.value = "'" + value
            else:
                target.value = value
        for column, width in sheet.get("column_widths", {}).items():
            if re.fullmatch(r"[A-Z]{1,3}", column):
                page.column_dimensions[column].width = min(80, max(5, float(width)))
        for chart in sheet.get("charts", [])[:10]:
            plot = LineChart()
            plot.title = str(chart.get("title", "Forecast"))
            plot.add_data(
                Reference(page, range_string=f"'{page.title}'!{chart['range']}"), titles_from_data=True
            )
            page.add_chart(plot, chart.get("anchor", "E2"))
    book.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
    book.save(path)


def presentation(path, spec, directory, store, sources=None):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slides = spec.get("slides", [])
    if not 1 <= len(slides) <= 100:
        raise ValueError("A deck requires 1–100 slides")
    previews, links, sources = [], [], sources or {}
    for index, item in enumerate(slides):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        title = str(item.get("title", ""))[:300]
        box = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(12), Inches(0.8))
        provenance.ppt_text(box.text_frame, title, sources, links)
        box.text_frame.paragraphs[0].font.size = Pt(32)
        layout = item.get("layout", "bullets")
        body = provenance.with_ids(item.get("body", ""), item)
        if layout in {"two-column", "quote", "title", "section", "bullets"}:
            columns = item.get("columns", [body]) if layout == "two-column" else [body]
            for column, text in enumerate(columns[:2]):
                shape = slide.shapes.add_textbox(
                    Inches(0.7 + column * 6.1),
                    Inches(1.7),
                    Inches(5.8 if len(columns) > 1 else 12),
                    Inches(4.8),
                )
                shape.text_frame.word_wrap = True
                provenance.ppt_text(shape.text_frame, text, sources, links)
        elif layout == "table":
            rows = item.get("rows", [])[:25]
            if not rows or not rows[0]:
                raise ValueError("A table slide needs rows")
            table = slide.shapes.add_table(
                len(rows), len(rows[0]), Inches(0.7), Inches(1.7), Inches(12), Inches(4.8)
            ).table
            for r, row in enumerate(rows):
                for c, value in enumerate(row[: len(rows[0])]):
                    provenance.ppt_text(
                        table.cell(r, c).text_frame, provenance.value_text(value), sources, links
                    )
        elif layout == "chart":
            chart = CategoryChartData()
            chart.categories = item.get("labels", [])[:50]
            for series in item.get("series", [])[:8]:
                chart.add_series(str(series["name"]), [float(value) for value in series["values"]][:50])
            slide.shapes.add_chart(
                XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.7), Inches(1.7), Inches(12), Inches(4.8), chart
            )
        elif layout == "image":
            images = store.rows(
                "SELECT filename,mime FROM attachments WHERE id=?", (item.get("attachment", ""),)
            )
            if not images or not images[0]["mime"].startswith("image/"):
                raise ValueError("Image slides require an existing image attachment")
            slide.shapes.add_picture(
                str(directory.parent / "attachments" / images[0]["filename"]),
                Inches(0.7),
                Inches(1.7),
                height=Inches(4.8),
            )
        else:
            raise ValueError("Unsupported slide layout")
        if layout in {"table", "chart", "image"} and item.get("source_ids"):
            caption = slide.shapes.add_textbox(Inches(0.7), Inches(6.35), Inches(12), Inches(0.4))
            provenance.ppt_text(caption.text_frame, provenance.with_ids("Sources:", item), sources, links)
        slide.notes_slide.notes_text_frame.text = str(item.get("notes", ""))[:5000]
        slide.shapes.add_picture(
            str(Path(__file__).parent / "web" / "brand" / "source-mark.png"),
            Inches(12.1),
            Inches(6.9),
            width=Inches(0.65),
        )
        previews.append(
            {
                "slide": index + 1,
                "title": title,
                "layout": layout,
                "body": body[:3000],
                "needs_review": bool(item.get("needs_review", False)),
                "notes": str(item.get("notes", ""))[:3000],
            }
        )
    previews.extend(provenance.ppt_appendix(deck, sources, links))
    deck.save(path)
    return {"slides": previews, "preview_kind": "structured"}


def document(path, title, spec, sources=None):
    from docx import Document
    from docx.shared import Inches

    doc, sources = Document(), sources or {}
    doc.add_heading(title, 0)
    if spec.get("executive_summary"):
        doc.add_heading("Executive summary", 1)
        provenance.doc_text(doc.add_paragraph(), spec["executive_summary"], sources)
    for section in spec.get("sections", [])[:100]:
        if section.get("heading"):
            doc.add_heading(str(section["heading"]), min(3, max(1, int(section.get("level", 1)))))
        provenance.doc_text(
            doc.add_paragraph(), provenance.with_ids(section.get("text", ""), section), sources
        )
        rows = section.get("table", [])[:100]
        if rows:
            table = doc.add_table(rows=0, cols=len(rows[0]))
            for row in rows:
                cells = table.add_row().cells
                for index, value in enumerate(row[: len(cells)]):
                    provenance.doc_text(cells[index].paragraphs[0], provenance.value_text(value), sources)
    provenance.doc_appendix(doc, sources)
    unverified = [str(value) for value in spec.get("sources", []) if not provenance.MARKER.search(str(value))]
    if unverified:
        doc.add_heading("Unverified references", 1)
        for reference in unverified:
            doc.add_paragraph(reference)
    doc.sections[0].footer.paragraphs[0].text = "OpenDuet"
    doc.sections[0].footer.paragraphs[0].add_run().add_picture(
        str(Path(__file__).parent / "web" / "brand" / "source-mark.png"), width=Inches(0.55)
    )
    doc.save(path)


def pdf(path, title, content, sources=None):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles, sources = getSampleStyleSheet(), sources or {}
    elements = [Paragraph(html.escape(title), styles["Title"])]
    for line in content.splitlines()[:3000]:
        elements.extend([Paragraph(provenance.pdf_text(line, sources), styles["BodyText"]), Spacer(1, 5)])
    if sources:
        elements.append(Paragraph("Sources", styles["Heading1"]))
    for source in sources.values():
        elements.append(
            Paragraph(
                '<a name="src_' + source["id"] + '"/>' + html.escape(source["label"]), styles["Heading3"]
            )
        )
        elements.append(Paragraph(html.escape(source["excerpt"]), styles["BodyText"]))

    def footer(canvas, document):
        canvas.drawImage(
            str(Path(__file__).parent / "web" / "brand" / "source-mark.png"),
            42,
            24,
            width=45,
            height=22.5,
            mask="auto",
        )
        canvas.setFont("Helvetica", 8)
        canvas.drawString(94, 30, "OpenDuet")

    SimpleDocTemplate(str(path), pagesize=A4, title=title).build(
        elements, onFirstPage=footer, onLaterPages=footer
    )


def dashboard(spec):
    if not isinstance(spec, dict):
        raise ValueError("A dashboard requires a JSON object")
    if set(spec) - {"tiles", "tables", "charts", "where_to_look", "sources", "refresh"}:
        raise ValueError("Dashboard contains unsupported fields")
    for collection in ("tiles", "tables", "charts", "where_to_look"):
        if not isinstance(spec.get(collection, []), list) or len(spec.get(collection, [])) > 100:
            raise ValueError("Dashboard collections must be bounded lists")
    for tile in spec.get("tiles", []):
        if not isinstance(tile, dict) or not {"label", "value"} <= tile.keys():
            raise ValueError("Each dashboard tile needs a label and a value")
    return spec


def dashboard_image(path, title, spec):
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1200, 800), "#f5f3ed")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.load_default(size=36)
    body_font = ImageFont.load_default(size=22)
    small_font = ImageFont.load_default(size=17)
    draw.text((48, 36), str(title)[:60], fill="#292e27", font=title_font)
    for index, tile in enumerate(spec.get("tiles", [])[:6]):
        x, y = 48 + (index % 3) * 375, 115 + (index // 3) * 125
        draw.rounded_rectangle((x, y, x + 350, y + 105), radius=12, fill="#fffef9", outline="#d9ddd1")
        draw.text((x + 18, y + 12), str(tile["label"])[:26], fill="#62695f", font=small_font)
        draw.text((x + 18, y + 44), str(tile["value"])[:20], fill="#292e27", font=body_font)
    top = 395 if len(spec.get("tiles", [])) > 3 else 265
    for table in spec.get("tables", [])[:1]:
        draw.text((48, top), str(table.get("title", ""))[:70], fill="#292e27", font=body_font)
        top += 40
        for row in [table.get("columns", []), *table.get("rows", [])][:8]:
            draw.text(
                (48, top), "  |  ".join(str(value)[:24] for value in row[:4]), fill="#292e27", font=small_font
            )
            top += 30
    for chart in spec.get("charts", [])[:1]:
        values = [float(value) for value in chart.get("values", [])[:8]]
        maximum = max([abs(value) for value in values] + [1])
        for index, value in enumerate(values):
            x = 55 + index * 125
            height = int(120 * max(0, value) / maximum)
            draw.rectangle((x, 700 - height, x + 90, 700), fill="#4e5c36")
            draw.text(
                (x, 708),
                str(chart.get("labels", list(range(len(values))))[index])[:10],
                fill="#292e27",
                font=small_font,
            )
    draw.text(
        (48, 764),
        "OpenDuet · " + time.strftime("%Y-%m-%d %H:%M UTC", time.gmtime()),
        fill="#62695f",
        font=small_font,
    )
    brand = Image.open(Path(__file__).parent / "web" / "brand" / "source-mark.png").convert("RGBA")
    brand.thumbnail((70, 35))
    image.paste(brand, (1080, 750), brand)
    image.save(path, format="PNG")


def slide_images(path, store, ident):
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    renderer = shutil.which("pdftoppm")
    if not executable or not renderer:
        return {}
    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary)
        converted = subprocess.run(
            [
                executable,
                "-env:UserInstallation=" + (root / "profile").as_uri(),
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(root),
                str(path),
            ],
            capture_output=True,
            timeout=60,
            check=False,
        )
        document = root / (path.stem + ".pdf")
        if converted.returncode or not document.is_file():
            return {}
        output = path.parent / "previews" / ident
        output.mkdir(parents=True, exist_ok=True, mode=0o700)
        rendered = subprocess.run(
            [
                renderer,
                "-png",
                "-scale-to",
                "1200",
                "-f",
                "1",
                "-l",
                "100",
                str(document),
                str(output / "slide"),
            ],
            capture_output=True,
            timeout=60,
            check=False,
        )
        images = sorted(output.glob("slide-*.png")) if not rendered.returncode else []
        return {
            "preview_kind": "rendered" if images else "structured",
            "slide_images": [image.name for image in images],
        }


def recalculate(path):
    """Optional real Office validation, in a fresh LibreOffice profile."""
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        return {"recalculated": False, "note": "LibreOffice not installed; formulas recalculate when opened"}
    with tempfile.TemporaryDirectory() as directory:
        profile = Path(directory) / "profile"
        output = Path(directory) / "output"
        output.mkdir()
        result = subprocess.run(
            [
                executable,
                "-env:UserInstallation=" + profile.as_uri(),
                "--headless",
                "--convert-to",
                "xlsx",
                "--outdir",
                str(output),
                str(path),
            ],
            capture_output=True,
            timeout=60,
            check=False,
        )
        converted = output / path.name
        if result.returncode or not converted.is_file():
            return {"recalculated": False, "note": "LibreOffice conversion failed; inspect the workbook"}
        from openpyxl import load_workbook

        book = load_workbook(converted, data_only=True, read_only=True)
        errors = [
            f"{page.title}!{cell.coordinate}: {cell.value}"
            for page in book
            for row in page
            for cell in row
            if cell.data_type == "e"
        ]
        book.close()
        if errors:
            raise ValueError("Spreadsheet formula errors: " + "; ".join(errors[:20]))
        shutil.copyfile(converted, path)
        return {"recalculated": True, "formula_errors": []}


def publish(store, directory, task, title, format, content, parent, refresh, plain):
    ensure(store)
    directory.mkdir(parents=True, exist_ok=True, mode=0o700)
    if not parent:
        scheduled = store.rows("SELECT root FROM dashboard_runs WHERE task=?", (task,))
        parent = scheduled[0]["root"] if scheduled else None
    if parent and not store.rows("SELECT id FROM artifacts WHERE id=?", (parent,)):
        raise ValueError("Parent artifact does not exist")
    if refresh:
        if format != "dashboard":
            raise ValueError("Only dashboards support scheduled refresh")
        ZoneInfo(refresh["timezone"])
        datetime.strptime(refresh["local_time"], "%H:%M")
        if not str(refresh.get("prompt", "")).strip():
            raise ValueError("A refresh needs a work prompt")
    # Exports carry the same sources as the task. The host supplies these, not the model.
    from .citations import Sources

    citations = Sources(store)
    citations.require(task, content)
    sources = citations.listing(task)
    lookup = {row["id"]: row for row in sources}
    used = set(provenance.MARKER.findall(content))
    spec = (
        json.loads(content)
        if content.lstrip().startswith("{") and format in {"dashboard", "webapp", "xlsx", "pptx", "docx"}
        else None
    )
    if spec is None and format == "docx":
        spec = {"sections": [{"text": content}]}
    elif spec is None and format == "pptx":
        spec = {
            "slides": [
                {"title": title, "layout": "bullets", "body": block} for block in content.split("\n---\n")
            ]
        }

    def check_source_fields(value):
        if isinstance(value, dict):
            if "source_ids" in value:
                ids = value["source_ids"]
                if (
                    not isinstance(ids, list)
                    or len(ids) > 20
                    or any(not isinstance(ident, str) or ident not in lookup for ident in ids)
                ):
                    raise ValueError("Artifact source_ids must refer to this task's retrieved sources")
                used.update(ids)
            if isinstance(value.get("source"), str) and value["source"] in lookup:
                used.add(value["source"])
            for child in value.values():
                check_source_fields(child)
        elif isinstance(value, list):
            for child in value:
                check_source_fields(child)

    if spec is not None:
        check_source_fields(spec)
    if used:
        lookup = {ident: row for ident, row in lookup.items() if ident in used}
    details = {ident: provenance.detail(row) for ident, row in lookup.items()}
    source_lines = [
        f"[^{row['id']}]: {row['label']} {row['url']} — {' '.join(row['excerpt'].split())}"
        for row in details.values()
    ]
    metadata, source_sheet = {}, "Sources"
    if source_lines:
        if spec is not None and format in {"dashboard", "docx"}:
            spec["sources"] = source_lines
            content = json.dumps(spec, ensure_ascii=False)
        elif spec is not None and format == "xlsx":
            names = {sheet.get("name", "").lower() for sheet in spec.get("sheets", [])}
            while source_sheet.lower() in names:
                source_sheet += "_"
            spec.setdefault("sheets", []).append(
                {
                    "name": source_sheet,
                    "cells": {
                        "A1": "Retrieved task sources",
                        **{f"A{i + 2}": value for i, value in enumerate(source_lines)},
                    },
                }
            )
            content = json.dumps(spec, ensure_ascii=False)
        elif spec is not None and format == "pptx":
            content = json.dumps(spec, ensure_ascii=False)
        elif format in {"md", "txt"}:
            content += "\n\nSources\n" + "\n".join(source_lines)
        elif format == "csv":
            extra = io.StringIO()
            csv.writer(extra).writerows(
                [
                    ["Sources"],
                    ["ID", "Source", "URL", "Retrieved excerpt"],
                    *[[row["id"], row["label"], row["url"], row["excerpt"]] for row in details.values()],
                ]
            )
            content = content.rstrip() + "\n" + extra.getvalue()
        elif format == "html":
            content += (
                "<section><h2>Sources</h2>"
                + "".join(
                    '<p id="src_'
                    + row["id"]
                    + '">'
                    + html.escape(row["label"])
                    + ": "
                    + html.escape(row["excerpt"])
                    + "</p>"
                    for row in details.values()
                )
                + "</section>"
            )
    if format not in {"dashboard", "webapp", "pdf"} and spec is None:
        item = plain(store, directory, task, title, format, content)
        ident = item["id"]
    else:
        ident = uid()
        extension = {"dashboard": "json", "webapp": "zip"}.get(format, format)
        filename = ident + "." + extension
        path = directory / filename
        if format == "dashboard":
            path.write_text(json.dumps(dashboard(spec), ensure_ascii=False))
        elif format == "webapp":
            files = spec.get("files", {}) if spec else {"index.html": content}
            if not files or "index.html" not in files or len(files) > 50:
                raise ValueError("Prototype requires index.html and at most 50 files")
            with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                for name, data in files.items():
                    if (
                        Path(name).is_absolute()
                        or ".." in Path(name).parts
                        or Path(name).suffix not in {".html", ".css", ".js", ".json"}
                    ):
                        raise ValueError("Invalid prototype file")
                    archive.writestr(name, str(data))
            preview_dir = directory.parent / "prototypes" / ident
            preview_dir.mkdir(parents=True, mode=0o700)
            for name, data in files.items():
                target = preview_dir / name
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_text(str(data))
            metadata = {"preview_kind": "sandboxed", "preview_id": ident}
        elif format == "xlsx":
            workbook(path, spec, lookup, source_sheet)
            metadata = recalculate(path)
        elif format == "pptx":
            metadata = presentation(path, spec, directory, store, details)
            metadata.update(slide_images(path, store, ident))
        elif format == "docx":
            document(path, title, spec, details)
        elif format == "pdf":
            pdf(path, title, content, details)
        else:
            raise ValueError("Unsupported structured artifact format")
        path.chmod(0o600)
        if format == "dashboard":
            attachment = uid()
            images = directory.parent / "attachments"
            images.mkdir(exist_ok=True, mode=0o700)
            preview = images / (attachment + ".png")
            dashboard_image(preview, title, spec)
            preview.chmod(0o600)
            store.execute(
                "INSERT INTO attachments VALUES(?,?,?,?,?,?)",
                (attachment, task, title + ".png", "image/png", preview.name, preview.stat().st_size),
            )
            metadata["dashboard_image"] = attachment
        store.execute(
            "INSERT INTO artifacts VALUES(?,?,?,?,?,?)", (ident, task, title, format, filename, time.time())
        )
        item = {"id": ident, "title": title, "format": format, "download": f"/api/artifacts/{ident}"}
    metadata["sources"] = list(details.values())
    root = ident
    if parent:
        original = store.rows("SELECT * FROM artifacts WHERE id=?", (parent,))
        if not original:
            raise ValueError("Parent artifact does not exist")
        existing = store.rows("SELECT root FROM artifact_versions WHERE id=?", (parent,))
        root = existing[0]["root"] if existing else parent
        if not existing:
            store.execute("INSERT INTO artifact_versions VALUES(?,?,1,NULL,'','{}')", (parent, parent))
    with store.db() as db:
        db.execute("BEGIN IMMEDIATE")
        version = db.execute(
            "SELECT COALESCE(MAX(version),0)+1 FROM artifact_versions WHERE root=?", (root,)
        ).fetchone()[0]
        db.execute(
            "INSERT INTO artifact_versions VALUES(?,?,?,?,?,?)",
            (ident, root, version, parent, content, json.dumps(metadata)),
        )
    if refresh:
        ZoneInfo(refresh["timezone"])
        datetime.strptime(refresh["local_time"], "%H:%M")
        provider = store.task(task)["provider"]
        store.execute(
            "INSERT OR REPLACE INTO dashboard_refresh VALUES(?,?,?,?,?)",
            (root, refresh["local_time"], refresh["timezone"], refresh["prompt"], provider),
        )
    return item | {"root": root, "version": version, "metadata": metadata}


class Make:
    def __init__(self, runtime):
        self.runtime, self.store = runtime, runtime.store
        ensure(self.store)

    def attach_parent_sources(self, parent, task):
        rows = self.store.rows(
            "SELECT a.task,v.spec,v.metadata FROM artifacts a JOIN artifact_versions v ON v.id=a.id WHERE a.id=?",
            (parent,),
        )
        if not rows:
            return
        row = rows[0]
        metadata = json.loads(row["metadata"])
        ids = {source["id"] for source in metadata.get("sources", [])}
        ids.update(provenance.MARKER.findall(row["spec"]))
        ids.update(re.findall(r'"(s[0-9a-f]{12})"', row["spec"]))
        selected = [source for source in self.runtime.sources.listing(row["task"]) if source["id"] in ids]
        if not selected:
            return
        with self.store.db() as db:
            for source in selected:
                db.execute(
                    "INSERT OR IGNORE INTO sources VALUES(?,?,?,?,?)",
                    (source["id"], task, source["tool"], source["url"], source["excerpt"]),
                )
            for file in db.execute(
                "SELECT a.id FROM attachments a JOIN attachment_inputs i ON i.id=a.id WHERE a.task=? OR EXISTS(SELECT 1 FROM task_attachment_access p WHERE p.task=? AND p.id=a.id)",
                (row["task"], row["task"]),
            ).fetchall():
                if any(file["id"] in source["excerpt"] for source in selected):
                    db.execute("INSERT OR IGNORE INTO task_attachment_access VALUES(?,?)", (task, file["id"]))
            db.execute(
                "UPDATE tasks SET request=request||? WHERE id=?",
                (
                    "\nUNTRUSTED parent artifact source snapshots, now available to this task via read_source: "
                    + json.dumps([source["id"] for source in selected]),
                    task,
                ),
            )

    def revise(self, ident, changes):
        rows = self.store.rows("SELECT * FROM artifacts WHERE id=?", (ident,))
        if not rows:
            raise ValueError("Artifact not found")
        original = next(
            (version.get("spec", "") for version in self.versions(ident) if version["id"] == ident), ""
        )
        prompt = f"Revise artifact {ident} and publish a new version with parent={ident}. User requested changes:\n{str(changes)[:6000]}\nUNTRUSTED previous deliverable:\n{original[:20000]}"
        original_task = self.store.task(rows[0]["task"])
        task = self.store.new_task(original_task["channel"], prompt, original_task["provider"])
        self.attach_parent_sources(ident, task)
        route = self.store.rows("SELECT * FROM task_surfaces WHERE task=?", (rows[0]["task"],))
        if route:
            self.store.execute(
                "INSERT INTO task_surfaces VALUES(?,'telegram',?,?,NULL,NULL)",
                (task, route[0]["chat"], route[0]["topic"]),
            )
        return task

    def versions(self, ident):
        rows = self.store.rows("SELECT root FROM artifact_versions WHERE id=?", (ident,))
        root = rows[0]["root"] if rows else ident
        return self.store.rows(
            "SELECT v.*,a.title,a.format,a.created FROM artifact_versions v JOIN artifacts a ON a.id=v.id WHERE root=? ORDER BY version DESC",
            (root,),
        )

    def diff(self, ident, other):
        rows = self.store.rows("SELECT id,spec FROM artifact_versions WHERE id IN (?,?)", (ident, other))
        contents = {row["id"]: row["spec"] for row in rows}
        if ident not in contents or other not in contents:
            raise ValueError("Artifact version not found")
        return "\n".join(
            difflib.unified_diff(
                contents[other].splitlines(), contents[ident].splitlines(), fromfile=other, tofile=ident
            )
        )

    def tick(self, now=None):
        for job in self.store.rows("SELECT * FROM dashboard_refresh"):
            local = datetime.fromtimestamp(now or time.time(), ZoneInfo(job["timezone"]))
            day = local.date().isoformat()
            if local.strftime("%H:%M") < job["local_time"]:
                continue
            with self.store.db() as db:
                db.execute("BEGIN IMMEDIATE")
                if db.execute(
                    "SELECT 1 FROM dashboard_runs WHERE root=? AND local_day=?", (job["root"], day)
                ).fetchone():
                    continue
                task = self.store.new_task(
                    "main",
                    job["prompt"] + "\nPublish a new artifact version with parent=" + job["root"],
                    job["provider"],
                    db,
                )
                db.execute("INSERT INTO dashboard_runs VALUES(?,?,?)", (job["root"], day, task))
                db.execute("INSERT INTO task_origins VALUES(?,'dashboard',?)", (task, job["root"]))
                owner = self.runtime.telegram.settings()
                if owner.get("paired"):
                    db.execute(
                        "INSERT INTO task_surfaces VALUES(?,'telegram',?,NULL,NULL,NULL)",
                        (task, owner["owner_chat"]),
                    )
