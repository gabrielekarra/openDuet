"""Real LibreOffice acceptance; fails when the optional native converters are unavailable."""

import asyncio
import io
import json
import shutil
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

from openduet.artifacts import create
from openduet.attachments import ingest
from openduet.config import Config
from openduet.runtime import Runtime
from openduet.store import Store


async def main():
    assert shutil.which("libreoffice") or shutil.which("soffice"), "Install LibreOffice"
    assert shutil.which("pdftoppm"), "Install poppler-utils"
    with tempfile.TemporaryDirectory() as tmp:
        directory = Path(tmp)
        runtime = Runtime(Config(data_dir=directory), Store(directory))
        task = runtime.store.new_task("main", "Office acceptance", "local")
        input_bytes, input_book = io.BytesIO(), Workbook()
        input_book.active.title = "Revenue"
        input_book.active["A1"] = 100
        input_book.save(input_bytes)
        uploaded = await ingest(runtime, input_bytes.getvalue(), "budget.xlsx", task)
        retrieved = json.loads(
            await runtime.tools.invoke(
                "read_attachment", {"attachment": uploaded["attachment"], "locator": "'Revenue'!A1"}, task
            )
        )
        source = retrieved["sources"][0]["source_id"]
        spec = {
            "sheets": [
                {
                    "name": "Assumptions",
                    "cells": {"A1": {"value": 100, "source_ids": [source]}, "A2": {"value": 0.1}},
                },
                {
                    "name": "Forecast",
                    "cells": {
                        **{
                            f"A{i}": {"formula": f"=Assumptions!$A$1*(1+Assumptions!$A$2)^{i}"}
                            for i in range(1, 13)
                        },
                        "A13": {"formula": "=SUM(A1:A12)"},
                    },
                },
            ]
        }
        original = create(runtime.store, directory / "artifacts", task, "Forecast", "xlsx", json.dumps(spec))
        spec["sheets"][0]["cells"]["A1"]["value"] = 200
        spec["sheets"][0]["cells"]["A1"].pop("source_ids")
        revised = create(
            runtime.store,
            directory / "artifacts",
            task,
            "Forecast revised",
            "xlsx",
            json.dumps(spec),
            parent=original["id"],
        )
        totals = []
        for item in [original, revised]:
            assert item["metadata"]["recalculated"], item
            book = load_workbook(directory / "artifacts" / (item["id"] + ".xlsx"), data_only=True)
            totals.append(book["Forecast"]["A13"].value)
            if item == original:
                assert source in book["Assumptions"]["A1"].comment.text
                assert book["Assumptions"]["A1"].hyperlink is not None
            book.close()
        assert abs(totals[1] - 2 * totals[0]) < 1e-6, totals
        deck = create(
            runtime.store,
            directory / "artifacts",
            task,
            "Deck",
            "pptx",
            json.dumps(
                {
                    "slides": [
                        {
                            "title": "Real slide",
                            "layout": "two-column",
                            "columns": ["Revenue", "Costs"],
                            "notes": "Review the assumptions",
                            "needs_review": True,
                            "source_ids": [source],
                        }
                    ]
                }
            ),
        )
        assert deck["metadata"]["preview_kind"] == "rendered", deck
        assert deck["metadata"]["slide_images"], deck
        await runtime.close()
    print(
        "Office acceptance passed: live 12-month formulas, assumption revision, preserved source links/comments, native PPTX source-appendix previews"
    )


asyncio.run(main())
