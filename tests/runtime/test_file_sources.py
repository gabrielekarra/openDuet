"""Precise inert input snapshots, task boundaries and editable export provenance."""

import io
import json
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from reportlab.pdfgen.canvas import Canvas

from openduet import artifacts, file_sources, make
from openduet.attachments import ingest
from openduet.config import Config
from openduet.runtime import Runtime
from openduet.store import Store


@pytest.fixture
async def runtime(tmp_path):
    runtime = Runtime(Config(data_dir=tmp_path), Store(tmp_path))
    yield runtime
    await runtime.close()


def sheet_bytes():
    buffer, book = io.BytesIO(), Workbook()
    book.active.title = "Revenue"
    book.active["A1"] = 100
    book.active["B1"] = "=A1*2"
    book.active["C1"] = '=WEBSERVICE("https://example.test/do-not-call")'
    book.save(buffer)
    return buffer.getvalue()


async def attached(runtime, data=None, name="budget.xlsx"):
    item = await ingest(runtime, data if data is not None else sheet_bytes(), name)
    task = runtime.store.new_task(
        "main", "Review this input.\nUNTRUSTED attachment: " + json.dumps(item), "local"
    )
    return task, item


async def test_exact_cells_are_inert_and_source_records_keep_file_locator(runtime):
    task, item = await attached(runtime)
    assert "'Revenue'!A1: 100" in item["text"]
    history = runtime.initial_history(runtime.store.task(task))
    assert "read_attachment" in history[-2]["content"]
    result = json.loads(
        await runtime.tools.invoke("read_attachment", {"attachment": item["attachment"]}, task)
    )
    records = result["data"]["items"]
    assert records[0]["locator"] == "'Revenue'!A1" and records[0]["value"] == 100
    assert records[1]["formula"] and records[1]["cached_value"] is None
    assert not records[1]["evaluated_by_host"]
    assert records[2]["value"].startswith("=WEBSERVICE")
    source = next(row for row in result["sources"] if row["locator"] == "'Revenue'!A1")
    stored = next(row for row in runtime.sources.listing(task) if row["id"] == source["source_id"])
    assert "budget.xlsx" in stored["excerpt"] and "Revenue" in stored["excerpt"]
    other = runtime.store.new_task("main", "Unrelated", "local")
    with pytest.raises(ValueError, match="this task"):
        await runtime.tools.invoke("read_attachment", {"attachment": item["attachment"]}, other)
    with pytest.raises(ValueError, match="current task"):
        await runtime.tools.invoke("read_source", {"source_id": source["source_id"]}, other)
    file = runtime.store.rows("SELECT filename FROM attachments WHERE id=?", (item["attachment"],))[0]
    (runtime.config.data_dir / "attachments" / file["filename"]).write_bytes(b"changed")
    with pytest.raises(ValueError, match="changed since import"):
        file_sources.read(runtime, task, item["attachment"])


def test_document_pdf_slide_and_csv_locators():
    buffer, document = io.BytesIO(), Document()
    document.add_paragraph("Revenue 100")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Total 200"
    document.save(buffer)
    records, _ = file_sources.extract(buffer.getvalue(), "report.docx")
    assert [(r["locator"], r["value"]) for r in records] == [
        ("paragraph 1", "Revenue 100"),
        ("table 1!R1C1", "Total 200"),
    ]
    buffer = io.BytesIO()
    pdf = Canvas(buffer)
    for index in (1, 2):
        pdf.drawString(40, 700, f"Page {index} revenue {index * 100}")
        pdf.showPage()
    pdf.save()
    records, _ = file_sources.extract(buffer.getvalue(), "report.pdf")
    assert [r["locator"] for r in records] == ["page 1", "page 2"]
    assert "200" in records[1]["value"]
    buffer, deck = io.BytesIO(), Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Sales 100"
    deck.save(buffer)
    records, _ = file_sources.extract(buffer.getvalue(), "deck.pptx")
    assert records[0]["locator"] == "slide 1, shape 1" and records[0]["value"] == "Sales 100"
    records, _ = file_sources.extract(b'Client,Amount\n"Multi\nline",100\n', "rows.csv")
    assert any(r["locator"] == "CSV!A2" and r["value"] == "Multi\nline" for r in records)


def test_formula_cache_is_read_from_uploaded_bytes_not_computed():
    from xml.etree import ElementTree

    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(sheet_bytes())) as original, zipfile.ZipFile(buffer, "w") as changed:
        for entry in original.infolist():
            data = original.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(data)
                value = root.find(
                    ".//s:c[@r='B1']/s:v", {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                )
                value.text = "777"
                data = ElementTree.tostring(root)
            changed.writestr(entry, data)
    records, _ = file_sources.extract(buffer.getvalue(), "cache.xlsx")
    formula = next(record for record in records if record["locator"] == "'Revenue'!B1")
    assert formula["value"] == "=A1*2"
    assert formula["cached_value"] == 777 and not formula["evaluated_by_host"]


async def test_file_read_pages_have_budget_and_no_silent_loss(runtime):
    data = ("x" * 4000 + "\n") * 25
    task, item = await attached(runtime, data.encode(), "long.txt")
    offset, locators = 0, []
    while offset is not None:
        page = file_sources.read(runtime, task, item["attachment"], offset=offset)
        assert len(json.dumps(page)) < 30000
        locators.extend(record["locator"] for record in page["items"])
        offset = page["next_offset"]
    assert locators == [f"line {i}" for i in range(1, 26)]


async def test_brief_keeps_each_event_link_and_retrieved_id(runtime):
    records = [
        {"id": "one", "text": "Review 100 EUR", "links": ["https://mail.google.com/mail/u/0/#all/one"]},
        {"id": "two", "text": "Review 200 EUR", "links": ["https://mail.google.com/mail/u/0/#all/two"]},
    ]
    runtime.attention.events.ingest("gmail", records)
    for row in runtime.store.rows("SELECT * FROM connector_events"):
        await runtime.attention.triage(row)
    task = await runtime.attention.brief(force=True)
    sources = runtime.sources.listing(task)
    assert {source["url"] for source in sources} == {record["links"][0] for record in records}
    assert all(source["id"] in runtime.store.task(task)["request"] for source in sources)


async def test_large_inputs_use_explicit_references_and_invalid_input_rolls_back(tmp_path, monkeypatch):
    import httpx

    from openduet.server import create_app

    token = "fixture-private-token-not-production"
    monkeypatch.setenv("OPENDUET_TOKEN", token)
    app = create_app(Config(data_dir=tmp_path))
    runtime = app.state.runtime
    try:
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app),
            base_url="http://testserver",
            headers={"Authorization": "Bearer " + token},
        ) as client:
            upload = await client.post("/api/attachments?name=large.txt", content=("line\n" * 10000).encode())
            assert upload.status_code == 201
            ident = upload.json()["attachment"]
            result = await client.post(
                "/api/messages", json={"content": "Review this long file", "attachments": [ident]}
            )
            assert result.status_code == 202
            task = result.json()["task"]
            assert len(runtime.store.task(task)["request"]) < 100
            assert file_sources.read(runtime, task, ident)["items"][0]["locator"] == "line 1"
            assert upload.json()["truncated"]
            count = len(runtime.store.rows("SELECT id FROM tasks"))
            denied = await client.post(
                "/api/messages", json={"content": "Invalid reference", "attachments": ["missing"]}
            )
            assert denied.status_code == 400
            assert len(runtime.store.rows("SELECT id FROM tasks")) == count
    finally:
        await runtime.close()


async def test_local_cell_sources_link_in_office_pdf_and_markdown_and_survive_revision(runtime, monkeypatch):
    from pypdf import PdfReader

    task, item = await attached(runtime)
    result = json.loads(
        await runtime.tools.invoke(
            "read_attachment", {"attachment": item["attachment"], "locator": "'Revenue'!A1"}, task
        )
    )
    ident = result["sources"][0]["source_id"]
    directory = runtime.config.data_dir / "artifacts"
    monkeypatch.setattr(make, "recalculate", lambda path: {})
    monkeypatch.setattr(make, "slide_images", lambda *args: {})
    specs = {
        "docx": {
            "sections": [
                {
                    "text": "Revenue 100",
                    "source_ids": [ident],
                    "table": [[{"value": 100, "source_ids": [ident]}]],
                }
            ]
        },
        "pptx": {
            "slides": [
                {
                    "title": "Revenue",
                    "layout": "chart",
                    "labels": ["Sales"],
                    "series": [{"name": "EUR", "values": [100]}],
                    "source_ids": [ident],
                }
            ]
        },
        "xlsx": {"sheets": [{"name": "Sources", "cells": {"A1": {"value": 100, "source_ids": [ident]}}}]},
        "pdf": "Revenue 100 [^" + ident + "]",
        "md": "Revenue 100 [^" + ident + "]",
    }
    outputs = {
        kind: artifacts.create(
            runtime.store,
            directory,
            task,
            "Sourced",
            kind,
            json.dumps(spec) if isinstance(spec, dict) else spec,
        )
        for kind, spec in specs.items()
    }
    with zipfile.ZipFile(directory / (outputs["docx"]["id"] + ".docx")) as archive:
        xml = archive.read("word/document.xml").decode()
        assert 'w:anchor="src_' + ident + '"' in xml and "budget.xlsx" in xml and "Revenue" in xml
    with zipfile.ZipFile(directory / (outputs["pptx"]["id"] + ".pptx")) as archive:
        assert b"ppaction://hlinksldjump" in archive.read("ppt/slides/slide1.xml")
        assert b"budget.xlsx" in archive.read("ppt/slides/slide2.xml")
    book = load_workbook(directory / (outputs["xlsx"]["id"] + ".xlsx"))
    assert book["Sources"]["A1"].hyperlink.location.startswith("'Sources_'!")
    assert "'Revenue'!A1" in book["Sources"]["A1"].comment.text
    book.close()
    reader = PdfReader(directory / (outputs["pdf"]["id"] + ".pdf"))
    assert any(page.get("/Annots") for page in reader.pages)
    assert "budget.xlsx" in "".join(page.extract_text() for page in reader.pages)
    assert "'Revenue'!A1" in (directory / (outputs["md"]["id"] + ".md")).read_text()
    revised = runtime.make.revise(outputs["docx"]["id"], "Keep the same sourced value")
    assert ident in [row["id"] for row in runtime.sources.listing(revised)]
    inherited = file_sources.read(runtime, revised, item["attachment"], locator="'Revenue'!A1")
    assert inherited["items"][0]["value"] == 100
    artifacts.create(
        runtime.store,
        directory,
        revised,
        "Revision",
        "docx",
        json.dumps(specs["docx"]),
        parent=outputs["docx"]["id"],
    )
    next_artifact = runtime.store.rows("SELECT id FROM artifacts WHERE task=?", (revised,))[0]["id"]
    third = runtime.make.revise(next_artifact, "One more revision")
    assert (
        file_sources.read(runtime, third, item["attachment"], locator="'Revenue'!A1")["items"][0]["value"]
        == 100
    )
    runtime.store.status(task, "ready")
    runtime.store.status(revised, "ready")
    runtime.store.status(third, "ready")
    runtime.store.delete_channel("main")
    assert not runtime.store.rows("SELECT * FROM attachment_extracts")
    assert not runtime.store.rows("SELECT * FROM task_attachment_access")
