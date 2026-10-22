import io
import json
import sys
import types
import zipfile
from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook

from openduet.artifacts import create
from openduet.config import Config
from openduet.runtime import Runtime
from openduet.store import Store
from openduet.tools import Tool, schema


@pytest.fixture
def runtime(tmp_path):
    return Runtime(Config(data_dir=tmp_path), Store(tmp_path))


async def test_mac_opening_is_installed_app_only_and_no_shell(runtime):
    calls = []
    runtime.mac.apps = lambda: [{"name": "Calculator", "bundle_id": "com.apple.calculator"}]

    async def command(*args):
        calls.append(args)

    runtime.mac.command = command
    assert await runtime.mac.open_app("com.apple.calculator") == {"opened": "com.apple.calculator"}
    assert calls == [("/usr/bin/open", "-b", "com.apple.calculator")]
    with pytest.raises(ValueError):
        await runtime.mac.open_app("unknown; rm -rf /")
    assert len(calls) == 1
    await runtime.close()


async def test_mac_action_is_snapshot_bound_and_one_use(runtime, monkeypatch):
    performed = []
    workspace = types.SimpleNamespace(
        frontmostApplication=lambda: types.SimpleNamespace(processIdentifier=lambda: 123)
    )
    monkeypatch.setitem(
        sys.modules,
        "AppKit",
        types.SimpleNamespace(NSWorkspace=types.SimpleNamespace(sharedWorkspace=lambda: workspace)),
    )
    monkeypatch.setitem(
        sys.modules,
        "ApplicationServices",
        types.SimpleNamespace(
            AXUIElementPerformAction=lambda element, action: performed.append((element, action)) or 0
        ),
    )
    runtime.mac.check = lambda: None
    runtime.mac.pid = 123
    runtime.mac.snapshot_id = "snapshot"
    runtime.mac.elements = {"1": "native-control"}
    assert await runtime.mac.act("snapshot", "1", "click") == {"performed": "click"}
    with pytest.raises(ValueError):
        await runtime.mac.act("snapshot", "1", "click")
    assert performed == [("native-control", "AXPress")]
    await runtime.close()


async def test_owner_screen_command_delivers_photo_without_chat_pixels(runtime):
    runtime.mac.check = lambda: None
    runtime.mac.observe = lambda: {"app": "Fixture", "elements": []}

    async def command(*args):
        Path(args[-1]).write_bytes(b"fixture-png-bytes")

    runtime.mac.command = command
    token = "123456:fixture-only-telegram-token"
    runtime.vault.put("telegram.token", token)
    runtime.telegram.save(connected=True, paired=True, owner_id=12, owner_chat=12)
    requests = []

    def handler(request):
        requests.append(request)
        return httpx.Response(200, json={"ok": True, "result": {"message_id": 4, "chat": {"id": 12}}})

    await runtime.telegram.client.aclose()
    runtime.telegram.client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    await runtime.telegram.receive(
        {
            "update_id": 1,
            "message": {
                "message_id": 1,
                "from": {"id": 12},
                "chat": {"id": 12, "type": "private"},
                "text": "/screen",
            },
        }
    )
    await runtime.telegram.flush_one()
    assert requests[0].url.path.endswith("/sendPhoto")
    assert b"fixture-png-bytes" in requests[0].content
    assert runtime.store.rows("SELECT * FROM messages") == []
    assert "fixture-png-bytes" not in json.dumps(runtime.store.rows("SELECT * FROM events"))
    await runtime.telegram.receive(
        {
            "update_id": 2,
            "message": {"message_id": 2, "from": {"id": 99}, "chat": {"id": 99}, "text": "/screen"},
        }
    )
    assert len(runtime.store.rows("SELECT * FROM attachments")) == 1
    await runtime.close()


async def test_edit_then_approve_exact_payload_single_execution(runtime):
    calls = []

    async def send(arguments, task):
        calls.append(arguments)
        return {"receipt": "sent-once"}

    runtime.tools.catalog["send_email"] = Tool(
        "send_email", "Send email", schema({"body": {"type": "string"}}), send, True
    )
    task = runtime.store.new_task("main", "Send a reviewed email", "local")
    history = [
        {
            "role": "assistant",
            "content": "",
            "tool_calls": [
                {
                    "id": "call",
                    "type": "function",
                    "function": {"name": "send_email", "arguments": '{"body": "original"}'},
                }
            ],
        }
    ]
    runtime.checkpoint(task, history, 0)
    assert not await runtime.process_calls(runtime.store.task(task), history)
    old = runtime.store.rows("SELECT * FROM calls")[0]
    replacement = runtime.approvals.decide(old["id"], "edit", {"body": "edited"})["id"]
    with pytest.raises(ValueError):
        runtime.approvals.decide(old["id"], "approve")
    runtime.approvals.decide(replacement, "approve")
    with pytest.raises(ValueError):
        runtime.approvals.decide(replacement, "approve")
    updated = json.loads(runtime.store.task(task)["history"])
    assert await runtime.process_calls(runtime.store.task(task), updated)
    assert calls == [{"body": "edited"}]
    await runtime.close()


async def test_deferral_allows_remaining_work(runtime):
    async def send(arguments, task):
        raise AssertionError("Deferred action executed")

    runtime.tools.catalog["send_email"] = Tool("send_email", "Send email", schema({}), send, True)
    task = runtime.store.new_task("main", "Prepare independent work", "local")
    history = [
        {
            "role": "assistant",
            "tool_calls": [{"id": "call", "function": {"name": "send_email", "arguments": "{}"}}],
        }
    ]
    runtime.checkpoint(task, history, 0)
    assert not await runtime.process_calls(runtime.store.task(task), history)
    call = runtime.store.rows("SELECT * FROM calls")[0]
    runtime.approvals.decide(call["id"], "defer")
    assert await runtime.process_calls(runtime.store.task(task), history)
    assert json.loads(history[-1]["content"])["deferred"] is True
    assert runtime.store.rows("SELECT status FROM calls")[0]["status"] == "deferred"
    await runtime.close()


async def test_financial_workbook_live_formulas_and_version_history(runtime):
    task = runtime.store.new_task("main", "Build a 12 month forecast", "local")
    cells = {"A1": {"value": "Revenue forecast"}}
    for month in range(1, 13):
        cells[f"A{month + 1}"] = {"value": month}
        cells[f"B{month + 1}"] = {"formula": f"=Assumptions!$A$1*(1+Assumptions!$A$2)^A{month + 1}"}
    cells["B14"] = {"formula": "=SUM(B2:B13)", "number_format": "0.00"}
    spec = {
        "sheets": [
            {"name": "Assumptions", "cells": {"A1": {"value": 100}, "A2": {"value": 0.1}}},
            {"name": "Forecast", "cells": cells},
        ]
    }
    item = create(
        runtime.store, runtime.config.data_dir / "artifacts", task, "Forecast", "xlsx", json.dumps(spec)
    )
    book = load_workbook(runtime.config.data_dir / "artifacts" / f"{item['id']}.xlsx", data_only=False)
    assert book["Forecast"]["B14"].value == "=SUM(B2:B13)"
    assert book["Forecast"]["B2"].data_type == "f"
    book.close()
    spec["sheets"][0]["cells"]["A1"]["value"] = 200
    revised = create(
        runtime.store,
        runtime.config.data_dir / "artifacts",
        task,
        "Forecast revised",
        "xlsx",
        json.dumps(spec),
        parent=item["id"],
    )
    assert revised["version"] == 2
    assert len(runtime.make.versions(revised["id"])) == 2
    assert (runtime.config.data_dir / "artifacts" / f"{item['id']}.xlsx").is_file()
    assert "100" in runtime.make.diff(revised["id"], item["id"])
    await runtime.close()


async def test_dashboard_refresh_uses_local_day_and_keeps_versions(runtime):
    from datetime import datetime
    from zoneinfo import ZoneInfo

    task = runtime.store.new_task("main", "Create dashboard", "local")
    item = create(
        runtime.store,
        runtime.config.data_dir / "artifacts",
        task,
        "Daily",
        "dashboard",
        json.dumps({"tiles": [{"label": "Revenue", "value": "€100", "source": "unverified"}]}),
        refresh={
            "local_time": "05:30",
            "timezone": "Europe/Rome",
            "prompt": "Refresh from connected accounts",
        },
    )
    stamp = datetime(2026, 3, 29, 5, 30, tzinfo=ZoneInfo("Europe/Rome")).timestamp()
    runtime.make.tick(stamp)
    runtime.make.tick(stamp + 60)
    runs = runtime.store.rows("SELECT * FROM dashboard_runs")
    assert len(runs) == 1
    assert runs[0]["root"] == item["id"]
    assert runtime.store.task(runs[0]["task"])["provider"] == "local"
    await runtime.close()


async def test_deck_review_notes_and_document_sources(runtime):
    task = runtime.store.new_task("main", "Prepare deck and brief", "local")
    item = create(
        runtime.store,
        runtime.config.data_dir / "artifacts",
        task,
        "Deck",
        "pptx",
        json.dumps(
            {
                "slides": [
                    {
                        "title": "Review this forecast",
                        "layout": "bullets",
                        "body": "Assumptions need review",
                        "notes": "Verify the source",
                        "needs_review": True,
                    }
                ]
            }
        ),
    )
    assert item["metadata"]["slides"][0]["needs_review"] is True
    from pptx import Presentation

    deck = Presentation(runtime.config.data_dir / "artifacts" / f"{item['id']}.pptx")
    assert "Verify the source" in deck.slides[0].notes_slide.notes_text_frame.text
    brief = create(
        runtime.store,
        runtime.config.data_dir / "artifacts",
        task,
        "Research",
        "docx",
        json.dumps(
            {
                "executive_summary": "Verified facts only",
                "sections": [
                    {
                        "heading": "Findings",
                        "text": "A sourced finding",
                        "table": [["Metric", "Value"], ["A", "1"]],
                    }
                ],
                "sources": ["https://example.test/source"],
            }
        ),
    )
    from docx import Document

    document = Document(runtime.config.data_dir / "artifacts" / f"{brief['id']}.docx")
    assert any("https://example.test/source" in p.text for p in document.paragraphs)
    await runtime.close()


async def test_plugin_static_install_and_traversal_rejection(runtime):
    def archive(files):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w") as output:
            for name, value in files.items():
                output.writestr(name, value)
        return stream.getvalue()

    good = archive(
        {
            ".claude-plugin/plugin.json": '{"name":"fixture","description":"A static fixture"}',
            "skills/fixture/SKILL.md": "---\nname: fixture\ndescription: fixture skill\n---\nUse verified sources.",
        }
    )
    runtime.plugins.install(good)
    assert runtime.plugins.load("fixture") == "Use verified sources."
    for data in (
        archive({"plugin.json": '{"name":"bad"}', "../escape.md": "bad"}),
        archive({"plugin.json": '{"name":"bad"}', "hooks/run.sh": "echo no"}),
    ):
        with pytest.raises(ValueError):
            runtime.plugins.install(data)
    runtime.plugins.uninstall("fixture")
    assert "fixture" not in runtime.plugins.skills()
    await runtime.close()


async def test_unknown_citation_is_flagged_for_text_tool_models(runtime):
    task = runtime.store.new_task("main", "Research", "local")
    registered = runtime.sources.register(task, "search", {"url": "https://example.test", "value": 12})
    output = runtime.sources.validate(task, f"12 [^{registered['source_id']}] and 99 [^invented]")
    assert "[^invented]" not in output and "Unverified source" in output
    assert registered["source_id"] in output
    await runtime.close()
