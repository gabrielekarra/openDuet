from __future__ import annotations

import asyncio
import json
import time
from types import SimpleNamespace

import httpx
import pytest

from openduet.artifacts import create
from openduet.browser import Browser
from openduet.config import Config, ConnectorConfig, ProviderConfig
from openduet.policy import check_payload
from openduet.providers import Provider, Reply
from openduet.runtime import Runtime
from openduet.server import create_app
from openduet.store import Store
from openduet.tools import Tool, schema


@pytest.fixture
def cfg(tmp_path, monkeypatch):
    monkeypatch.setenv("OPENDUET_TOKEN", "test-only-access-token-32-characters")
    return Config(data_dir=tmp_path)


def tool_reply(*calls):
    return Reply(
        {
            "role": "assistant",
            "content": "",
            "tool_calls": [
                {"id": str(i), "type": "function", "function": {"name": n, "arguments": json.dumps(a)}}
                for i, (n, a) in enumerate(calls)
            ],
        },
        10,
        10,
    )


def fake_factory(*replies):
    queue = list(replies)

    class Fake:
        def __init__(self, config):
            pass

        async def complete(self, messages, tools):
            assert messages[0]["role"] == "system"
            return queue.pop(0)

    return Fake


async def run(runtime, ident):
    runtime.store.status(ident, "running")
    await runtime.run_task(ident)


@pytest.mark.parametrize("kind", ["openai", "anthropic", "ollama"])
async def test_wire_adapters_and_secret_bindings(kind, monkeypatch):
    monkeypatch.setenv("TEST_PROVIDER_KEY", "a-private-value")
    received = []

    def request(req):
        body = json.loads(req.content)
        received.append(body)
        assert body["model"] == "test-model"
        if kind == "anthropic":
            assert req.headers["x-api-key"] == "a-private-value"
            assert req.url.path == "/v1/messages"
            return httpx.Response(
                200,
                json={
                    "content": [{"type": "tool_use", "id": "x", "name": "search", "input": {"q": "test"}}],
                    "usage": {"input_tokens": 7, "output_tokens": 3},
                },
            )
        if kind == "ollama":
            assert req.url.path == "/api/chat"
            assert body["stream"] is False
            return httpx.Response(
                200,
                json={
                    "message": {
                        "role": "assistant",
                        "content": "",
                        "tool_calls": [{"function": {"name": "search", "arguments": {"q": "test"}}}],
                    }
                },
            )
        assert req.headers["authorization"] == "Bearer a-private-value"
        return httpx.Response(
            200,
            json={
                "choices": [{"message": {"role": "assistant", "content": "Hello"}}],
                "usage": {"prompt_tokens": 7, "completion_tokens": 3},
            },
        )

    cfg = ProviderConfig(
        kind=kind,
        model="test-model",
        api_key_env="TEST_PROVIDER_KEY",
        base_url="https://model.example" + ("/v1" if kind != "ollama" else ""),
    )
    answer = await Provider(cfg, httpx.MockTransport(request)).complete(
        [{"role": "system", "content": "system"}, {"role": "user", "content": "hello"}], []
    )
    assert len(received) == 1
    assert "a-private-value" not in json.dumps(received)
    assert answer.input_tokens > 0 and answer.output_tokens > 0
    if kind != "openai":
        assert json.loads(answer.message["tool_calls"][0]["function"]["arguments"]) == {"q": "test"}


async def test_text_tool_protocol_and_no_auth_for_local():
    def request(req):
        assert "authorization" not in req.headers
        body = json.loads(req.content)
        assert "tools" not in body
        assert "Tool calling protocol" in body["messages"][0]["content"]
        return httpx.Response(
            200,
            json={
                "choices": [
                    {
                        "message": {
                            "role": "assistant",
                            "content": '```json\n{"tool_calls":[{"name":"search_context","arguments":{"query":"meeting"}}]}\n```',
                        }
                    }
                ]
            },
        )

    result = await Provider(ProviderConfig(tools=False), httpx.MockTransport(request)).complete(
        [{"role": "system", "content": "system"}, {"role": "user", "content": "hello"}], []
    )
    assert result.message["tool_calls"][0]["function"]["name"] == "search_context"


async def test_provider_error_does_not_leak_response():
    def request(req):
        return httpx.Response(401, text="private token and prompt echoed")

    with pytest.raises(RuntimeError, match="HTTP 401") as exc:
        await Provider(ProviderConfig(), httpx.MockTransport(request)).complete([], [])
    assert "private" not in str(exc.value)


async def test_review_then_execute_exactly_once_and_user_completes(cfg):
    app = create_app(
        cfg,
        start_worker=False,
        provider_factory=fake_factory(
            tool_reply(("remember", {"content": "Prefers morning meetings"})),
            Reply({"role": "assistant", "content": "Saved your preference."}, 5, 5),
        ),
    )
    runtime, store = app.state.runtime, app.state.store
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app),
        base_url="http://test",
        headers={"Authorization": "Bearer " + cfg.token()},
    ) as client:
        ident = (await client.post("/api/messages", json={"content": "Remember my preference"})).json()[
            "task"
        ]
        await run(runtime, ident)
        assert store.task(ident)["status"] == "awaiting_approval"
        assert not store.rows("SELECT * FROM memories")
        approval = (await client.get("/api/state")).json()["approvals"][0]
        assert json.loads(approval["arguments"])["content"] == "Prefers morning meetings"
        assert (
            await client.post(f"/api/approvals/{approval['id']}", json={"approve": True})
        ).status_code == 200
        assert (
            await client.post(f"/api/approvals/{approval['id']}", json={"approve": True})
        ).status_code == 409
        await run(runtime, ident)
        assert len(store.rows("SELECT * FROM memories")) == 1
        assert store.task(ident)["status"] == "ready"
        assert (await client.post(f"/api/tasks/{ident}/complete")).status_code == 200
        assert store.task(ident)["status"] == "completed"


async def test_safe_preparation_continues_while_approval_pending(cfg):
    runtime = Runtime(
        cfg,
        Store(cfg.data_dir),
        fake_factory(
            tool_reply(
                ("remember", {"content": "Needs approval"}),
                ("create_artifact", {"title": "Draft", "format": "md", "content": "Prepared for review"}),
            )
        ),
    )
    ident = runtime.store.new_task("main", "Prepare", "local")
    await run(runtime, ident)
    assert runtime.store.task(ident)["status"] == "awaiting_approval"
    assert len(runtime.store.rows("SELECT * FROM artifacts")) == 1
    assert not runtime.store.rows("SELECT * FROM memories")


async def test_revocation_after_approval_blocks_effect(cfg):
    runtime = Runtime(
        cfg,
        Store(cfg.data_dir),
        fake_factory(
            tool_reply(("send_mail", {"text": "Hello"})),
            Reply({"role": "assistant", "content": "Access revoked"}),
        ),
    )
    effects = []

    async def send(arguments, task):
        effects.append(arguments)
        return {"sent": True}

    runtime.tools.catalog["send_mail"] = Tool(
        "send_mail", "send", schema({"text": {"type": "string"}}), send, True, "mail"
    )
    ident = runtime.store.new_task("main", "Send", "local")
    await run(runtime, ident)
    runtime.store.execute("UPDATE calls SET status='approved'")
    runtime.store.set_preference("disabled_connectors", ["mail"])
    await run(runtime, ident)
    assert not effects
    assert runtime.store.rows("SELECT status FROM calls")[0]["status"] == "denied"


async def test_uncertain_effect_is_never_replayed_after_restart(cfg):
    runtime = Runtime(cfg, Store(cfg.data_dir), fake_factory(tool_reply(("effect", {}))))
    effects = []

    async def effect(arguments, task):
        effects.append(1)
        raise TimeoutError("secret response")

    runtime.tools.catalog["effect"] = Tool("effect", "effect", schema({}), effect)
    ident = runtime.store.new_task("main", "Run effect", "local")
    await run(runtime, ident)
    assert runtime.store.task(ident)["status"] == "blocked"
    runtime.store.recover()
    await run(runtime, ident)
    assert effects == [1]
    assert runtime.store.rows("SELECT status FROM calls")[0]["status"] == "uncertain"
    assert "secret" not in runtime.store.task(ident)["error"]


async def test_cancel_marks_inflight_effect_uncertain(cfg):
    started = asyncio.Event()
    runtime = Runtime(cfg, Store(cfg.data_dir), fake_factory(tool_reply(("effect", {}))))

    async def effect(arguments, task):
        started.set()
        await asyncio.Event().wait()

    runtime.tools.catalog["effect"] = Tool("effect", "effect", schema({}), effect)
    ident = runtime.store.new_task("main", "Run", "local")
    task = asyncio.create_task(run(runtime, ident))
    runtime.current = (ident, task)
    await started.wait()
    runtime.cancel(ident)
    with pytest.raises(asyncio.CancelledError):
        await task
    assert runtime.store.task(ident)["status"] == "cancelled"
    assert runtime.store.rows("SELECT status FROM calls")[0]["status"] == "uncertain"


async def test_malformed_tool_args_do_not_execute(cfg):
    runtime = Runtime(
        cfg,
        Store(cfg.data_dir),
        fake_factory(
            tool_reply(("remember", {"content": "x", "extra": "unauthorized"})),
            Reply({"role": "assistant", "content": "Invalid request"}),
        ),
    )
    ident = runtime.store.new_task("main", "Remember", "local")
    await run(runtime, ident)
    assert not runtime.store.rows("SELECT * FROM memories")
    assert runtime.store.rows("SELECT status FROM calls")[0]["status"] == "failed"


async def test_budget_blocks_model_call(cfg):
    cfg.daily_token_budget = 1
    runtime = Runtime(cfg, Store(cfg.data_dir), fake_factory())
    runtime.store.usage(2, 2)
    ident = runtime.store.new_task("main", "Hello", "local")
    await run(runtime, ident)
    assert runtime.store.task(ident)["status"] == "blocked"
    assert "budget" in runtime.store.task(ident)["error"]


def test_schedule_claim_is_durable_and_does_not_overlap(cfg):
    store = Store(cfg.data_dir)
    runtime = Runtime(cfg, store)
    store.execute(
        "INSERT INTO schedules VALUES('schedule','Daily','Brief','main','local',3600,?,1)", (time.time() - 1,)
    )
    runtime.tick_schedules()
    runtime.tick_schedules()
    assert len(store.rows("SELECT * FROM tasks")) == 1
    store.execute("UPDATE schedules SET next_run=0")
    runtime.tick_schedules()
    assert len(store.rows("SELECT * FROM tasks")) == 1
    store.execute("UPDATE tasks SET status='completed'")
    runtime.tick_schedules()
    assert len(store.rows("SELECT * FROM tasks")) == 2


def test_search_and_channel_deletion_remove_context(cfg):
    store = Store(cfg.data_dir)
    store.execute("INSERT INTO channels VALUES('side','Side project')")
    ident = store.new_task("side", "Invoice nebula", "local")
    store.status(ident, "ready")
    assert store.search("nebula")[0]["task"] == ident
    store.set_preference("search_chats", False)
    assert store.search("nebula") == []
    store.set_preference("search_chats", True)
    store.delete_channel("side")
    assert store.search("nebula") == []
    assert not store.rows("SELECT * FROM tasks")
    assert store.search('" AND invoice --') == []


@pytest.mark.parametrize("format", ["md", "html", "csv", "xlsx", "docx", "pptx"])
def test_generated_artifacts_are_real_formats(cfg, format):
    store = Store(cfg.data_dir)
    item = create(
        store,
        cfg.data_dir / "artifacts",
        "task",
        "Test report",
        format,
        'Revenue,100\n=HYPERLINK("https://example.com"),200\n---\nNext slide\nConclusion',
    )
    path = cfg.data_dir / "artifacts" / (item["id"] + "." + format)
    assert path.is_file()
    if format == "xlsx":
        from openpyxl import load_workbook

        assert load_workbook(path).active["A2"].data_type == "s"
    elif format == "docx":
        from docx import Document

        assert Document(path).paragraphs[0].text == "Test report"
    elif format == "pptx":
        from pptx import Presentation

        assert len(Presentation(path).slides) == 2


async def test_auth_csrf_cookie_and_artifact_isolation(cfg):
    app = create_app(cfg, start_worker=False)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app), base_url="http://test") as client:
        assert (await client.get("/api/state")).status_code == 401
        assert (
            await client.post(
                "/api/login", json={"token": cfg.token()}, headers={"Origin": "http://attacker.example"}
            )
        ).status_code == 403
        signed_in = await client.post("/api/login", json={"token": cfg.token()})
        assert signed_in.status_code == 200
        assert "HttpOnly" in signed_in.headers["set-cookie"]
        assert cfg.token() not in signed_in.headers["set-cookie"]
        assert (await client.get("/api/state")).status_code == 200
        item = create(
            app.state.store, cfg.data_dir / "artifacts", "task", "evil", "html", "<script>alert(1)</script>"
        )
        response = await client.get(item["download"])
        assert "attachment" in response.headers["content-disposition"]
        assert response.headers["content-type"] == "application/octet-stream"
        assert "no-store" in response.headers["cache-control"]
        assert (await client.post("/api/logout")).status_code == 200
        assert (await client.get("/api/state")).status_code == 401


def test_browser_allowlist_and_config_rejection(cfg):
    browser = Browser(cfg.data_dir, ["example.com"])
    assert browser.allowed("https://example.com/path")
    for url in [
        "https://example.com.evil.org",
        "file:///etc/passwd",
        "http://example.com",
        "https://user:pass@example.com",
    ]:
        assert not browser.allowed(url)
    with pytest.raises(ValueError):
        ConnectorConfig(url="http://untrusted.example/mcp")
    with pytest.raises(ValueError):
        ProviderConfig(base_url="https://user:secret@example.com")


def test_secrets_never_enter_prompt_or_receipt(cfg, monkeypatch):
    monkeypatch.setenv("MODEL_KEY", "secret-value-from-environment")
    cfg.providers["local"].api_key_env = "MODEL_KEY"
    runtime = Runtime(cfg, Store(cfg.data_dir))
    runtime.private_values.add("private-login-value")
    assert runtime.redact("secret-value-from-environment private-login-value") == "[REDACTED] [REDACTED]"


def test_sensitive_identifiers_are_rejected_before_execution():
    with pytest.raises(ValueError, match="payment card"):
        check_payload({"body": "4111 1111 1111 1111"})
    with pytest.raises(ValueError, match="identification"):
        check_payload({"body": "123-45-6789"})
    check_payload({"body": "Invoice INV-1234, amount 200"})


async def test_reused_call_id_cannot_change_approved_payload(cfg):
    first = tool_reply(("remember", {"content": "first action"}))
    different = tool_reply(("remember", {"content": "different action"}))
    runtime = Runtime(cfg, Store(cfg.data_dir), fake_factory(first, different))
    ident = runtime.store.new_task("main", "Remember", "local")
    await run(runtime, ident)
    runtime.store.execute("UPDATE calls SET status='approved'")
    await run(runtime, ident)
    assert runtime.store.task(ident)["status"] == "failed"
    assert [m["content"] for m in runtime.store.rows("SELECT * FROM memories")] == ["first action"]


async def test_private_login_requires_the_page_user_reviewed(cfg):
    browser = Browser(cfg.data_dir, ["example.com", "other.example"])
    browser.context = object()
    browser.page = SimpleNamespace(url="https://other.example/login")
    with pytest.raises(ValueError, match="page changed"):
        await browser.login("input", "private value", "https://example.com/login")


async def test_private_login_errors_do_not_echo_credentials(cfg, monkeypatch):
    cfg.browser = True
    app = create_app(cfg, start_worker=False)

    async def failing_login(*args):
        raise ValueError("fill(private-password) failed")

    monkeypatch.setattr(app.state.runtime.browser, "login", failing_login)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app),
        base_url="http://test",
        headers={"Authorization": "Bearer " + cfg.token()},
    ) as client:
        result = await client.post(
            "/api/browser/login",
            json={
                "selector": "input",
                "value": "private-password",
                "expected_url": "https://example.com/login",
            },
        )
        assert result.status_code == 400
        assert "private-password" not in result.text
